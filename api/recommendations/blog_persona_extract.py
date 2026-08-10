"""Text extraction for blog-persona file uploads (see
migrations/013_blog_personas.sql, migrations/014_blog_personas_categories.sql
for the fields this feeds). A person's source material commonly already
exists as a spreadsheet of survey stats or a written buyer-persona/
testimonial doc, not something they want to retype - this lets them upload
that file directly instead of pasting.

Deliberately narrow: only .md/.txt (already plain text - covers buyer
persona and testimonials, which is what's actually being written by hand)
and .xlsx (the actual non-text format in use, for stats/survey exports). No
.docx/.pdf - not asked for, and PDF extraction quality (especially scanned
PDFs needing OCR) is a much bigger lift than this warrants.
"""

import io
import re

from openpyxl import load_workbook

_TEXT_EXTENSIONS = (".md", ".txt")

# Markdown exported from some editors (Notion, Obsidian, Word-to-Markdown,
# a browser's "copy as Markdown") embeds pasted images inline as base64 data
# URIs rather than linking an external file - either `![alt](data:...)` or
# reference-style `[label]: <data:...>`. That's not persona/testimonial
# content, it's a screenshot someone pasted while writing the doc, and it
# can be hundreds of KB of base64 noise per image - dumping it into the
# textarea as-is buries the actual text. Stripped to a short placeholder
# instead, regardless of which markdown syntax wraps it.
_DATA_URI_RE = re.compile(r"data:[\w.+-]+/[\w.+-]+;base64,[A-Za-z0-9+/=]+")

# Markdown exported from Notion/Obsidian/Word etc. also tends to carry
# formatting cruft that's meaningless once dropped into a plain textarea:
# blockquote markers on every line of a quoted testimonial (`> like this`,
# or nested `>> like this`), non-breaking spaces in place of regular ones,
# trailing whitespace (often a markdown hard-line-break, i.e. two trailing
# spaces - irrelevant since nothing here renders markdown), and runs of
# several blank lines from a converted list/table. Scoped to the .md/.txt
# branch only (see extract_text) - a spreadsheet cell legitimately starting
# with ">" (e.g. an ">50%" survey answer) must NOT be mistaken for a
# blockquote marker, so this never touches .xlsx-derived text.
_BLOCKQUOTE_RE = re.compile(r"^[ \t]*>+[ \t]?", re.MULTILINE)
_TRAILING_WS_RE = re.compile(r"[ \t]+$", re.MULTILINE)
_EXCESS_BLANK_LINES_RE = re.compile(r"\n{3,}")
_NBSP = " "


class UnsupportedFileType(ValueError):
    pass


def _strip_embedded_media(text):
    return _DATA_URI_RE.sub("[embedded image omitted]", text)


def _normalize_markdown_noise(text):
    text = text.replace(_NBSP, " ")
    text = _BLOCKQUOTE_RE.sub("", text)
    text = _TRAILING_WS_RE.sub("", text)
    text = _EXCESS_BLANK_LINES_RE.sub("\n\n", text)
    return text.strip()


def extract_text(filename, raw_bytes):
    """Best-effort plain-text extraction from an uploaded file, keyed off
    its extension. Raises UnsupportedFileType for anything else."""
    name = (filename or "").lower()
    if name.endswith(_TEXT_EXTENSIONS):
        text = raw_bytes.decode("utf-8", errors="replace").strip()
        text = _normalize_markdown_noise(text)
    elif name.endswith(".xlsx"):
        text = _extract_xlsx(raw_bytes)
    else:
        raise UnsupportedFileType(f"Unsupported file type: {filename!r}. Upload a .md, .txt, or .xlsx file.")
    return _strip_embedded_media(text)


def _extract_xlsx(raw_bytes):
    """Renders every sheet as tab-separated rows. Blank rows are preserved
    as blank lines - spreadsheet survey exports are usually grouped into
    question/answer blocks separated by blank rows, and keeping those blank
    lines keeps that grouping legible as plain text (matches the shape of
    the tab-separated survey data this was built against). Sheet-name
    headers are only added when there's more than one sheet, so a
    single-sheet workbook - the common case - round-trips as plain rows
    with no extra noise."""
    workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=True)
    multi_sheet = len(workbook.sheetnames) > 1
    lines = []
    for sheet in workbook.worksheets:
        if multi_sheet:
            if lines:
                lines.append("")
            lines.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            lines.append("\t".join(cells))
    return "\n".join(lines).strip()
